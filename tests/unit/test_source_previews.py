"""The cited page of a knowledge-base document: converted to PDF, found from
quotes, rendered with marks, and served through both contrib layers (tests/urls.py)."""

from io import BytesIO

import pytest
from django.core.files.base import ContentFile

pytest.importorskip("pdf_oxide")
pytest.importorskip("PIL")

from PIL import Image  # noqa: E402

from django_ai_sdk.files.previews import (  # noqa: E402
    PAGE_DPI,
    passage_areas,
    quote_page,
    render_page,
    to_pdf,
)

REPEATED = "These rules apply to every employee."
SHARED = "Holiday hours are paid out in May."
ALSO = "Night shifts earn a bonus of twenty percent."


def make_pdf(pages: list[list[str]]) -> bytes:
    """A PDF with one page per item: its lines in Helvetica, from the top."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [%s] /Count %d >>"
        % (b" ".join(b"%d 0 R" % (4 + 2 * i) for i in range(len(pages))), len(pages)),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, lines in enumerate(pages):
        text = b"".join(b"(%s) Tj T* " % line.encode() for line in lines)
        stream = b"BT /F1 12 Tf 14 TL 72 720 Td " + text + b"ET"
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (5 + 2 * i)
        )
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    pdf, offsets = b"%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    pdf += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return pdf


PDF = make_pdf(
    [
        [REPEATED, "Chapter one."],
        [REPEATED, "Holiday hours are paid out", "in May.", ALSO],
        [REPEATED, SHARED],
    ]
)


class TestQuotePage:
    def test_the_page_the_passages_share_wins_over_a_repeated_sentence(self):
        assert quote_page(PDF, [REPEATED, SHARED, ALSO]) == 2

    def test_a_quote_broken_over_lines_is_found(self):
        assert quote_page(PDF, [ALSO, SHARED]) == 2

    def test_the_first_page_on_a_tie(self):
        assert quote_page(PDF, [REPEATED]) == 1

    def test_too_short_quotes_find_nothing(self):
        assert quote_page(PDF, ["May.", "bonus"]) is None

    def test_only_the_first_quotes_count(self):
        assert (
            quote_page(PDF, [f"Nothing like this number {n}." for n in range(8)] + [ALSO]) is None
        )


class TestPassageAreas:
    @pytest.fixture
    def words(self):
        import pdf_oxide

        with pdf_oxide.PdfDocument.from_bytes(PDF) as pdf:
            return pdf.extract_words(1)

    def test_a_box_per_word_of_the_passage(self, words):
        assert len(passage_areas(words, "paid out in May")) == 4

    def test_sentence_by_sentence_when_the_whole_is_not_printed(self, words):
        quote = f"This sentence is not on the page. {ALSO}"
        assert len(passage_areas(words, quote)) == len(ALSO.split())

    def test_nothing_for_text_not_on_the_page(self, words):
        assert passage_areas(words, "Not printed anywhere at all.") == []


WEEKEND = "Weekend work is paid at one and a half."
# ALSO twice: on the first line, and again right after WEEKEND
REPEATS = make_pdf([[ALSO, "Unrelated filler text sits right here.", WEEKEND, ALSO]])
FIRST_LINE = 720  # y of the first line, in points from the bottom


class TestPassageFoundOnce:
    """Located as the citation grounder locates it: once, not at every repeat."""

    @pytest.fixture
    def words(self):
        import pdf_oxide

        with pdf_oxide.PdfDocument.from_bytes(REPEATS) as pdf:
            return pdf.extract_words(0)

    def test_a_repeated_passage_is_marked_once(self, words):
        assert len(passage_areas(words, ALSO)) == len(ALSO.split())

    def test_a_slightly_different_passage_is_still_marked(self, words):
        assert passage_areas(words, ALSO.replace("twenty", "thirty"))

    def test_pieces_are_found_in_order_not_at_an_earlier_repeat(self, words):
        quote = f"{WEEKEND} This sentence is nowhere on the page at all. {ALSO}"

        boxes = passage_areas(words, quote)

        assert len(boxes) == len(WEEKEND.split()) + len(ALSO.split())
        assert all(y < FIRST_LINE for _, y, _, _ in boxes)


class TestRenderPage:
    def test_a_png_of_the_page_with_where_the_marks_are(self):
        png, mark = render_page(PDF, 2, [ALSO])

        image = Image.open(BytesIO(png))
        assert image.format == "PNG"
        assert image.width == round(612 * PAGE_DPI / 72)
        top, left, bottom, right = map(float, mark.split(","))
        assert 0 < top < bottom < 0.5  # near the top of the page
        assert 0 < left < right < 1

    def test_no_mark_without_quotes(self):
        assert render_page(PDF, 1, [])[1] == ""

    def test_no_such_page(self):
        assert render_page(PDF, 4, []) is None
        assert render_page(PDF, 0, []) is None


def png() -> bytes:
    out = BytesIO()
    Image.new("RGB", (40, 20), "white").save(out, "PNG")
    return out.getvalue()


def docx() -> bytes:
    import docx as python_docx

    document = python_docx.Document()
    for line in ["Chapter one.", *[f"Filler paragraph number {n}." for n in range(60)], ALSO]:
        document.add_paragraph(line)
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["Rule", "Value"])
    workbook.active.append([SHARED, "yes"])
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


class TestToPdf:
    @pytest.mark.parametrize(
        ("extension", "data"),
        [
            ("docx", docx),
            ("DOCX", docx),
            ("xlsx", xlsx),
            ("csv", lambda: f'Rule,Value\n"{SHARED}",yes\n'.encode()),
            ("md", lambda: f"# Rules\n\n{SHARED}\n".encode()),
            ("txt", lambda: SHARED.encode()),
        ],
    )
    def test_a_document_becomes_a_pdf_with_its_text_findable(self, extension, data):
        pdf = to_pdf(data(), extension)

        page = quote_page(pdf, [SHARED, ALSO])
        assert page is not None
        assert render_page(pdf, page, [SHARED, ALSO])[1]

    def test_a_long_document_finds_the_later_page(self):
        pdf = to_pdf(docx(), "docx")

        assert quote_page(pdf, [ALSO]) > 1

    def test_a_pdf_stays_as_it_is(self):
        assert to_pdf(PDF, ".pdf") is PDF

    def test_an_image_is_not_made_a_pdf(self):
        # shown as itself: no text, so nothing to mark on a page
        assert to_pdf(png(), "png") is None

    def test_no_preview_for_other_files(self):
        assert to_pdf(b"PK", "zip") is None


# Endpoints


@pytest.fixture
def owner(django_user_model):
    return django_user_model.objects.create(email="owner@example.com")


@pytest.fixture
def stranger(django_user_model):
    return django_user_model.objects.create(email="stranger@example.com")


@pytest.fixture
def stored(owner, settings, tmp_path):
    """A private knowledge base the owner may read: a PDF, a Word file, a zip and
    an image."""
    from django_ai_sdk.memories.models import Entry, EntryDocument
    from tests.factories.db import MemoryFactory, MemoryUserFactory

    settings.MEDIA_ROOT = str(tmp_path)
    memory = MemoryFactory(private=True)
    MemoryUserFactory(memory=memory, user=owner)

    def upload(name, data, content_type=""):
        entry = Entry.objects.create(memory=memory, name=name)
        doc = EntryDocument(
            memory=memory,
            entry=entry,
            file_name=name,
            file_extension=name.rsplit(".")[-1],
            content_type=content_type,
        )
        doc.file.save(name, ContentFile(data), save=True)
        return str(entry.id), str(doc.id)

    return (
        upload("rules.pdf", PDF),
        upload("rules.DOCX", docx()),
        upload("notes.zip", b"PK"),
        upload("chart.png", png(), "image/png"),
    )


BASES = ["/api/memories", "/drf/memories"]


def at(base: str, path: str) -> str:
    """``path`` on a contrib layer; DRF's routes end with a slash."""
    return f"{base}{path}/" if base.startswith("/drf") else f"{base}{path}"


@pytest.mark.django_db
@pytest.mark.parametrize("base", BASES)
class TestEndpoints:
    def test_the_document_and_the_page_of_the_quotes(self, client, owner, stored, base):
        (entry_id, doc_id), *_ = stored
        client.force_login(owner)

        response = client.get(
            at(base, "/source/document"), {"entry_id": entry_id, "quote": [SHARED, ALSO]}
        )

        assert response.status_code == 200
        found = response.json()
        url = at(base, f"/source/document/{doc_id}/file")
        assert found == {
            "id": doc_id,
            "name": "rules.pdf",
            "url": url,
            "page": 2,
            "page_url": f"{url}#page=2",
        }

    def test_the_pdf_opens_inline(self, client, owner, stored, base):
        (_, doc_id), *_ = stored
        client.force_login(owner)

        response = client.get(at(base, f"/source/document/{doc_id}/file"))

        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert response["Content-Disposition"].startswith("inline")
        assert response["X-Content-Type-Options"] == "nosniff"
        assert b"".join(response.streaming_content) == PDF

    def test_the_page_with_its_marks(self, client, owner, stored, base):
        (_, doc_id), *_ = stored
        client.force_login(owner)

        response = client.get(at(base, f"/source/document/{doc_id}/page/2"), {"quote": ALSO})

        assert response.status_code == 200
        assert response["Content-Type"] == "image/png"
        assert len(response["X-Mark"].split(",")) == 4

    def test_no_such_page(self, client, owner, stored, base):
        (_, doc_id), *_ = stored
        client.force_login(owner)

        response = client.get(at(base, f"/source/document/{doc_id}/page/9"))

        assert response.status_code == 404

    def test_not_for_files_without_a_preview(self, client, owner, stored, base):
        _, _, (entry_id, doc_id), _ = stored
        client.force_login(owner)

        assert client.get(at(base, "/source/document"), {"entry_id": entry_id}).status_code == 404
        assert client.get(at(base, f"/source/document/{doc_id}/file")).status_code == 404

    def test_not_for_someone_who_may_not_read_the_memory(self, client, stranger, stored, base):
        (entry_id, doc_id), *_ = stored
        client.force_login(stranger)

        for response in (
            client.get(at(base, "/source/document"), {"entry_id": entry_id}),
            client.get(at(base, f"/source/document/{doc_id}/file")),
            client.get(at(base, f"/source/document/{doc_id}/page/1")),
        ):
            assert response.status_code in (403, 404)

    def test_a_word_file_is_found_and_marked_but_downloads_as_itself(
        self, client, owner, stored, base
    ):
        from django_ai_sdk.memories.models import EntryDocument

        _, (entry_id, doc_id), *_ = stored
        client.force_login(owner)

        found = client.get(at(base, "/source/document"), {"entry_id": entry_id, "quote": ALSO})
        original = client.get(at(base, f"/source/document/{doc_id}/file"))
        page = client.get(
            at(base, f"/source/document/{doc_id}/page/{found.json()['page']}"), {"quote": ALSO}
        )

        assert found.json()["name"] == "rules.DOCX"
        assert found.json()["page"] > 1
        # #page=N only means something in a PDF viewer
        assert found.json()["page_url"] is None
        assert original["Content-Disposition"].startswith("attachment")
        assert 'filename="rules.DOCX"' in original["Content-Disposition"]
        assert original["Content-Type"] == "application/octet-stream"
        assert (
            b"".join(original.streaming_content) == EntryDocument.objects.get(id=doc_id).file.read()
        )
        assert page["Content-Type"] == "image/png"
        assert page["X-Mark"]

    def test_an_image_is_shown_as_itself(self, client, owner, stored, base):
        *_, (entry_id, doc_id) = stored
        client.force_login(owner)

        found = client.get(at(base, "/source/document"), {"entry_id": entry_id, "quote": ALSO})
        original = client.get(found.json()["url"])

        assert found.status_code == 200
        assert (found.json()["page"], found.json()["page_url"]) == (None, None)
        assert original["Content-Type"] == "image/png"
        assert "attachment" not in original.get("Content-Disposition", "")
        assert original["Content-Security-Policy"] == "sandbox"
        assert client.get(at(base, f"/source/document/{doc_id}/page/1")).status_code == 404
